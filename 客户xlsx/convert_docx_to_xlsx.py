import docx
import re
from openpyxl import Workbook

doc = docx.Document(r'd:\code空间\客户xlsx\店铺粉丝量完整清单-2026.docx')
wb = Workbook()
ws = wb.active
ws.title = '店铺粉丝量清单'
ws.append(['店铺', '粉丝量'])

seen = set()
noise_keywords = ['经过分析', '我已完整', '需要我', '图片内',
                  '非常抱歉', '漏了很多', '整理如下']

# Keywords after fan values that should be truncated
fan_cut_keywords = ['需要我', '要不要', '我帮你', '需要我把',
                    '需要我帮你', '我可以帮你', '需要吗', '要不要我',
                    '注：部分店铺', '注：所有店铺']

for p in doc.paragraphs:
    text = p.text.strip()
    if not text or len(text) < 20:
        continue
    
    entries = re.split(r'\d+[.、\)）]\s*', text)
    
    for entry in entries:
        entry = entry.strip()
        if not entry or len(entry) < 3:
            continue
        
        match = re.match(r'^(.+?)\s*( - | – | — |：|:|\s*[-—–]\s*)\s*(.*)$', entry)
        if not match:
            continue
        
        shop = match.group(1).strip()
        fans = match.group(3).strip()
        
        shop = re.sub(r'^###?\s*', '', shop).strip()
        shop = shop.strip('-—– ').strip()
        
        if any(kw in shop for kw in noise_keywords):
            continue
        if len(shop) < 2:
            continue
        
        # Skip table-formatted entries
        if '|' in shop:
            continue
        
        # Clean fan value
        fans = re.sub(r'[（(][^）)]*[）)]', '', fans)
        fans = fans.replace('粉丝', '').replace('**', '').strip()
        fans = re.sub(r'\s+', ' ', fans)
        
        # Truncate at common AI suggestion text
        for kw in fan_cut_keywords:
            idx = fans.find(kw)
            if idx >= 0:
                fans = fans[:idx].strip()
                break
        
        # If fans contains multiple entries (concatenated), keep only first fan value
        fans = re.sub(r'^(\d+(?:\.\d+)?万).*$', r'\1', fans)
        
        key = shop.lower().replace(' ', '').replace('-', '')
        if key in seen:
            continue
        seen.add(key)
        
        ws.append([shop, fans])

output_path = r'd:\code空间\客户xlsx\店铺粉丝量完整清单-2026.xlsx'
wb.save(output_path)
print(f'成功导出 {len(seen)} 条数据到: {output_path}')